"""Convert the old Google Sheet into seed/import.sql and data/weekly.json.

Usage:
  1. Download the sheet: File > Download > Microsoft Excel (.xlsx)
  2. python tools/import_sheet.py path/to/sheet.xlsx
  3. npm run db:import:local   (or db:import:remote)

Every form submission becomes one entry per filled-in challenge. Rows whose timestamp was
removed in the sheet (marked DELETED / CHEAT) are skipped.
"""
import json, re, sys
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
GUILD = "358698681672925195"
# Column ranges of the "Form responses" tab, by set.
SET_COLUMNS = {"original": range(2, 31), "expansion": range(31, 63), "mix": range(63, 99)}
PROOF_COLUMN = 99
# Pre-form scores have no timestamp; date them just before the first form response.
LEGACY_DATE = "2024-07-01T00:00:00.000Z"
# Final Mix scores above 1000 came from a since-fixed game bug.
BUGGED_SQL = "UPDATE entries SET bugged = 1 WHERE challenge_id = 'mix4_mix' AND score > 1000 AND source = 'sheet';"


# Misspelled names in the sheet -> the player's real name.
NAME_FIXES = {"Wawrzienic": "Wawrzyniec", "Shagrath": "Shagrat"}


def clean_name(name):
    name = re.sub(r"\s+", " ", str(name)).strip()
    return NAME_FIXES.get(name, name)


def normalize_proof(url):
    if not url:
        return None
    url = str(url).strip()
    m = re.match(r"^https?://(?:www\.|ptb\.|canary\.)?discord(?:app)?\.com/channels/(\d+)/(\d+)/(\d+)", url)
    if m:
        return f"https://discord.com/channels/{m[1]}/{m[2]}/{m[3]}"
    return url if url.startswith("https://") else None


def sql(value):
    if value is None:
        return "NULL"
    if isinstance(value, int):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def name_key(name):
    return re.sub(r"[^a-z]", "", name.lower())


def import_entries(wb, challenges):
    ws = wb["Form responses"]
    header = [c.value for c in ws[1]]
    by_set_name = {(c["set"], c["name"]): c["id"] for c in challenges}
    column_ids = {}
    for set_id, cols in SET_COLUMNS.items():
        for col in cols:
            column_ids[col] = by_set_name[(set_id, header[col])]

    lines, skipped = [], 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        timestamp, name = row[0], row[1]
        if timestamp is None or not name:
            skipped += 1
            continue
        player = clean_name(name)
        proof = normalize_proof(row[PROOF_COLUMN])
        created = timestamp.strftime("%Y-%m-%dT%H:%M:%S.000Z")
        for col, challenge_id in column_ids.items():
            value = row[col]
            if value is None:
                continue
            values = [challenge_id, player, player.lower(), int(round(value)), proof, created, "sheet"]
            lines.append(
                "INSERT INTO entries (challenge_id, player, player_key, score, proof_url, created_at, source) "
                f"VALUES ({', '.join(sql(v) for v in values)});"
            )
    # Before the form existed, a few players' scores were typed straight into the "Data" tab
    # (rows between the "Highscore" row and the first form response). Same column layout.
    first_form_name = next(r[1] for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None)
    for row in wb["Data"].iter_rows(min_row=12, values_only=True):
        if row[1] == first_form_name:
            break
        if not row[1]:
            continue
        player = clean_name(row[1])
        for col, challenge_id in column_ids.items():
            value = row[col]
            if not isinstance(value, (int, float)) or value <= 0:
                continue
            values = [challenge_id, player, player.lower(), int(round(value)), None, LEGACY_DATE, "sheet"]
            lines.append(
                "INSERT INTO entries (challenge_id, player, player_key, score, proof_url, created_at, source) "
                f"VALUES ({', '.join(sql(v) for v in values)});"
            )

    out = ROOT / "seed" / "import.sql"
    out.parent.mkdir(exist_ok=True)
    # Same rule as migrations/0002_bugged.sql, so a re-import keeps the bug flags.
    sql_text = "DELETE FROM entries WHERE source = 'sheet';\n" + "\n".join(lines) + "\n" + BUGGED_SQL + "\n"
    out.write_text(sql_text, encoding="utf-8")
    print(f"wrote {len(lines)} entries to {out} (skipped {skipped} deleted rows)")


def parse_results(text):
    """'Lemmingsplayer - 650', 'Wawrzienic - 382/wolvs - 332' or a tie 'mwi & wolvs - 309' -> [{name, score}]."""
    results = []
    for part in str(text or "").split("/"):
        m = re.match(r"^\s*(.+?)\s*-\s*(\d+)\s*$", part)
        if m:
            for name in re.split(r"\s*(?:&|,)\s*", m[1]):
                results.append({"name": clean_name(name), "score": int(m[2])})
    return results


def import_weekly(wb, challenges):
    ws = wb["Challenge Results"]
    # The weekly challenges were all from the Random Mix set; the sheet spells some names loosely.
    mix = {name_key(c["name"]): c["id"] for c in challenges if c["set"] == "mix"}
    aliases = {"finalmix": "mix4_mix", "rennaissancetimetravelers": "time_traveller2_mix"}
    weeks = []
    for row in ws.iter_rows(min_row=9, values_only=True):
        date, number, name, winner, runner_up = row[1:6]
        if not name:
            continue
        key = name_key(name)
        challenge_id = mix.get(key) or aliases.get(key)
        if not challenge_id:
            raise SystemExit(f"Unknown weekly challenge: {name!r}")
        weeks.append({
            "number": int(number),
            "date": date.strftime("%Y-%m-%d"),
            "challengeId": challenge_id,
            "winners": parse_results(winner),
            "runnersUp": parse_results(runner_up),
        })
    out = ROOT / "data" / "weekly.json"
    out.write_text(json.dumps(weeks, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"wrote {len(weeks)} weekly challenges to {out}")


def main(path):
    challenges = json.loads((ROOT / "data" / "challenges.json").read_text(encoding="utf-8"))
    wb = openpyxl.load_workbook(path, data_only=True)
    import_entries(wb, challenges)
    import_weekly(wb, challenges)


if __name__ == "__main__":
    main(sys.argv[1])
